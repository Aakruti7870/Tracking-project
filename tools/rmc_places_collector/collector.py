#!/usr/bin/env python3
"""Four-state Google Places RMC directory collector.

This is an isolated batch job. It does not import the TrackMyRMC application,
connect to the production database, or expose the Places key. The only writes
are local output files, a private Cloud Storage checkpoint, and final objects.
"""

from __future__ import annotations

import asyncio
import csv
import gzip
import hashlib
import json
import math
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PLACES_SEARCH_URL = "https://places.googleapis.com/v1/places:searchText"
PLACES_DETAILS_URL = "https://places.googleapis.com/v1/places/{place_id}"
TARGET_STATES = ("Maharashtra", "Goa", "Karnataka", "Gujarat")
SEARCH_TERMS = (
    "RMC Plant",
    "Ready Mix Concrete",
    "RMC Supplier",
    "Concrete Supplier",
    "Ready Mix Concrete Plant",
    "Ready Mixed Concrete",
    "RMC Concrete",
    "RMC batching plant",
    "Concrete batching plant",
    "Ready Mix supplier",
)

SEARCH_FIELDS = ",".join(
    (
        "places.id",
        "places.displayName",
        "places.formattedAddress",
        "places.addressComponents",
        "places.location",
        "places.googleMapsUri",
        "places.businessStatus",
        "places.primaryType",
        "places.types",
        "nextPageToken",
    )
)
CENTER_FIELDS = ",".join(
    (
        "places.id",
        "places.displayName",
        "places.formattedAddress",
        "places.addressComponents",
        "places.location",
        "places.types",
    )
)
DETAIL_FIELDS = ",".join(
    (
        "id",
        "displayName",
        "formattedAddress",
        "addressComponents",
        "location",
        "nationalPhoneNumber",
        "internationalPhoneNumber",
        "googleMapsUri",
        "websiteUri",
        "rating",
        "userRatingCount",
        "businessStatus",
        "primaryType",
        "primaryTypeDisplayName",
        "types",
    )
)

COVERAGE_SOURCE_URLS = {
    "LGD export 2026-06-21": (
        "https://lgdirectory.gov.in/downloadDirectory.do | "
        "https://gist.github.com/planemad/b2195c7feb506f8436659f36da1e58af"
    ),
    "Karnataka taluk supplement 2026": (
        "https://en.wikipedia.org/wiki/List_of_taluks_of_Karnataka | "
        "https://igod.gov.in/sg/KA/E042/organizations"
    ),
    "Administrative supplement": (
        "https://mumbaicity.gov.in/ | https://mumbaisuburban.gov.in/"
    ),
    "Adaptive high-density grid": "Google Places geographic grid",
}

HIGH_DENSITY_DISTRICTS = {
    ("Maharashtra", "Mumbai City"),
    ("Maharashtra", "Mumbai Suburban"),
    ("Maharashtra", "Thane"),
    ("Maharashtra", "Palghar"),
    ("Maharashtra", "Pune"),
    ("Maharashtra", "Raigad"),
    ("Maharashtra", "Nashik"),
    ("Maharashtra", "Nagpur"),
    ("Goa", "North Goa"),
    ("Goa", "South Goa"),
    ("Karnataka", "Bengaluru Urban"),
    ("Karnataka", "Bengaluru Rural"),
    ("Karnataka", "Bengaluru South"),
    ("Karnataka", "Mysuru"),
    ("Karnataka", "Dakshina Kannada"),
    ("Karnataka", "Belagavi"),
    ("Gujarat", "Ahmedabad"),
    ("Gujarat", "Surat"),
    ("Gujarat", "Vadodara"),
    ("Gujarat", "Rajkot"),
    ("Gujarat", "Gandhinagar"),
    ("Gujarat", "Kachchh"),
}

STATE_ALIASES = {
    "maharashtra": "Maharashtra",
    "goa": "Goa",
    "karnataka": "Karnataka",
    "gujarat": "Gujarat",
    "gujrat": "Gujarat",
}

DISTRICT_ALIASES = {
    "ahmednagar": "Ahilyanagar",
    "ahilya nagar": "Ahilyanagar",
    "aurangabad": "Chhatrapati Sambhajinagar",
    "osmanabad": "Dharashiv",
    "mumbai": "Mumbai City",
    "mumbai suburban district": "Mumbai Suburban",
    "bombay suburban": "Mumbai Suburban",
    "bangalore urban": "Bengaluru Urban",
    "bangalore rural": "Bengaluru Rural",
    "ramanagara": "Bengaluru South",
    "ramanagaram": "Bengaluru South",
    "bangalore south": "Bengaluru South",
    "bellary": "Ballari",
    "belgaum": "Belagavi",
    "bijapur": "Vijayapura",
    "gulbarga": "Kalaburagi",
    "mysore": "Mysuru",
    "ahmadabad": "Ahmedabad",
    "banas kantha": "Banaskantha",
    "chhotaudepur": "Chhota Udepur",
    "dohad": "Dahod",
    "kutch": "Kachchh",
    "kachchh district": "Kachchh",
    "mahesana": "Mehsana",
    "panch mahals": "Panchmahal",
    "sabar kantha": "Sabarkantha",
}

STRONG_RMC_RE = re.compile(
    r"\b(r\.?m\.?c\.?|ready[ -]?mix(?:ed)?\s+concrete|ready[ -]?mix\s+supplier|"
    r"concrete\s+batch(?:ing)?\s+plant|batch(?:ing)?\s+plant|concrete\s+plant)\b",
    re.IGNORECASE,
)
CONCRETE_CUE_RE = re.compile(r"\b(concrete|rmc|ready[ -]?mix)\b", re.IGNORECASE)
FALSE_POSITIVE_RE = re.compile(
    r"\b(tile|tiles|hardware|sanitary|plywood|paint|marble|granite|"
    r"cement\s+(shop|dealer|agency|store)|paver|paving\s+block|"
    r"concrete\s+pipe|precast\s+pipe|hume\s+pipe|block\s+manufacturer|"
    r"civil\s+contractor|construction\s+company|builder|architect)\b",
    re.IGNORECASE,
)

MASTER_COLUMNS = [
    "Sr No",
    "Plant / Supplier Name",
    "Record Type",
    "Public Business Phone",
    "International Phone",
    "Full Address",
    "State",
    "District",
    "Taluka / Tehsil",
    "City / Locality",
    "PIN Code",
    "Latitude",
    "Longitude",
    "Google Place ID",
    "Google Maps URL",
    "Website",
    "Rating",
    "Review Count",
    "Business Status",
    "Google Primary Type",
    "Search Term",
    "Data Quality",
    "Verification Status",
    "Last Verified Date",
    "Phone Status",
    "Discovery Queries",
    "Address Components JSON",
    "Searched Districts",
    "Searched Talukas",
]

COVERAGE_COLUMNS = [
    "State",
    "District",
    "Taluka",
    "Coverage Source",
    "Coverage Source URL",
    "Source Code",
    "Location Note",
    "Search Terms Executed",
    "Raw Results",
    "Unique Results",
    "Retained",
    "Excluded",
    "Search Requests",
    "Center Latitude",
    "Center Longitude",
    "Center Source",
    "Search Completion Status",
    "Errors",
]


class BudgetExceeded(RuntimeError):
    """The configured safety cap was reached before another billable call."""


class PlacesRequestError(RuntimeError):
    def __init__(self, status_code: int, status: str, message: str):
        super().__init__(f"Places request failed: {status_code} {status}")
        self.status_code = status_code
        self.status = status
        self.safe_message = message[:300]


@dataclass(frozen=True)
class Config:
    api_key: str
    output_bucket: str | None
    output_prefix: str
    local_output_dir: Path
    states: tuple[str, ...]
    max_text_requests: int
    max_detail_requests: int
    search_radius_m: int
    grid_radius_m: int
    request_delay_ms: int
    resume: bool
    max_pages: int
    checkpoint_every_units: int

    @classmethod
    def from_env(cls) -> "Config":
        key = next(
            (
                os.environ.get(name, "").strip()
                for name in (
                    "GOOGLE_PLACES_API_KEY",
                    "GOOGLE_MAPS_KEY",
                    "GOOGLE_MAPS_API_KEY",
                )
                if os.environ.get(name, "").strip()
            ),
            "",
        )
        requested = os.environ.get("STATES", "|".join(TARGET_STATES))
        states: list[str] = []
        for token in re.split(r"[|,]", requested):
            canonical = STATE_ALIASES.get(normalize_text(token))
            if canonical and canonical not in states:
                states.append(canonical)
        if not states:
            raise ValueError("STATES does not contain a supported state")
        return cls(
            api_key=key,
            output_bucket=os.environ.get("OUTPUT_BUCKET", "").strip() or None,
            output_prefix=os.environ.get("OUTPUT_PREFIX", "rmc-places").strip("/")
            or "rmc-places",
            local_output_dir=Path(
                os.environ.get("LOCAL_OUTPUT_DIR", "/tmp/rmc-places-output")
            ),
            states=tuple(states),
            max_text_requests=positive_int("MAX_TEXT_REQUESTS", 34_000),
            max_detail_requests=positive_int("MAX_DETAIL_REQUESTS", 6_500),
            search_radius_m=bounded_int("SEARCH_RADIUS_METERS", 35_000, 5_000, 50_000),
            grid_radius_m=bounded_int("GRID_RADIUS_METERS", 28_000, 5_000, 50_000),
            request_delay_ms=bounded_int("REQUEST_DELAY_MS", 100, 0, 5_000),
            resume=env_bool("RESUME", True),
            max_pages=bounded_int("MAX_PAGES_PER_QUERY", 3, 1, 3),
            checkpoint_every_units=bounded_int("CHECKPOINT_EVERY_UNITS", 10, 1, 100),
        )

    def validate(self) -> None:
        if not self.api_key:
            raise ValueError(
                "Google Places key is not configured. Inject GOOGLE_PLACES_API_KEY "
                "from Secret Manager."
            )
        if (
            not self.output_bucket
            and str(self.local_output_dir) == "/tmp/rmc-places-output"
        ):
            raise ValueError("OUTPUT_BUCKET is required for a Cloud Run execution")


def positive_int(name: str, default: int) -> int:
    try:
        value = int(os.environ.get(name, str(default)))
    except ValueError as exc:
        raise ValueError(f"{name} must be an integer") from exc
    if value <= 0:
        raise ValueError(f"{name} must be positive")
    return value


def bounded_int(name: str, default: int, low: int, high: int) -> int:
    value = (
        positive_int(name, default)
        if low > 0
        else int(os.environ.get(name, str(default)))
    )
    if not low <= value <= high:
        raise ValueError(f"{name} must be between {low} and {high}")
    return value


def env_bool(name: str, default: bool) -> bool:
    value = os.environ.get(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


def normalize_text(value: Any) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).split())


def normalize_dense(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def normalize_phone(value: Any) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) > 10 and digits.startswith("91"):
        digits = digits[2:]
    return digits[-10:] if len(digits) >= 10 else digits


def canonical_state(value: Any) -> str | None:
    return STATE_ALIASES.get(normalize_text(value))


def canonical_district(value: Any) -> str | None:
    raw = " ".join(str(value or "").strip().split())
    if not raw:
        return None
    normalized = normalize_text(raw.removesuffix(" district").removesuffix(" District"))
    alias = DISTRICT_ALIASES.get(normalized)
    if alias:
        return alias
    return " ".join(
        word.capitalize() if len(word) > 3 else word.title() for word in raw.split()
    )


def canonical_taluka(value: Any) -> str | None:
    raw = " ".join(str(value or "").strip().split())
    if not raw:
        return None
    cleaned = re.sub(
        r"\s+(taluka|taluk|tehsil|tahsil|sub[ -]?district)$",
        "",
        raw,
        flags=re.IGNORECASE,
    ).strip()
    return cleaned or raw


def component(place: dict[str, Any], *types: str) -> str | None:
    wanted = set(types)
    for item in place.get("addressComponents") or []:
        if wanted.intersection(item.get("types") or []):
            value = item.get("longText") or item.get("shortText")
            if value:
                return str(value).strip()
    return None


def display_name(place: dict[str, Any]) -> str:
    display = place.get("displayName") or {}
    return str(display.get("text") or "").strip()


def place_coordinates(place: dict[str, Any]) -> tuple[float | None, float | None]:
    location = place.get("location") or {}
    try:
        lat = float(location["latitude"])
        lng = float(location["longitude"])
    except (KeyError, TypeError, ValueError):
        return None, None
    if not (-90 <= lat <= 90 and -180 <= lng <= 180):
        return None, None
    return lat, lng


def haversine_m(a: tuple[float, float], b: tuple[float, float]) -> float:
    radius = 6_371_000.0
    lat1, lon1, lat2, lon2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    dlat, dlon = lat2 - lat1, lon2 - lon1
    h = (
        math.sin(dlat / 2) ** 2
        + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    )
    return 2 * radius * math.asin(math.sqrt(h))


class PlacesClient:
    def __init__(self, config: Config, counters: dict[str, int] | None = None):
        self.config = config
        self.counters = counters or {"text": 0, "details": 0}
        self.http = httpx.AsyncClient(timeout=httpx.Timeout(30.0, connect=10.0))

    async def close(self) -> None:
        await self.http.aclose()

    async def _sleep(self, attempt: int | None = None) -> None:
        delay = self.config.request_delay_ms / 1000
        if attempt is not None:
            delay += min(8.0, 0.5 * (2**attempt))
        if delay:
            await asyncio.sleep(delay)

    async def _request(
        self,
        method: str,
        url: str,
        *,
        field_mask: str,
        operation: str,
        json_body: dict[str, Any] | None = None,
        params: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        budget_key = "text" if operation == "text" else "details"
        cap = (
            self.config.max_text_requests
            if budget_key == "text"
            else self.config.max_detail_requests
        )
        headers = {
            "X-Goog-Api-Key": self.config.api_key,
            "X-Goog-FieldMask": field_mask,
        }
        if method == "POST":
            headers["Content-Type"] = "application/json"
        last_error: PlacesRequestError | None = None
        for attempt in range(5):
            if self.counters.get(budget_key, 0) >= cap:
                raise BudgetExceeded(f"{budget_key.upper()}_REQUEST_CAP_REACHED")
            # Count every HTTP attempt because retries can also be billable.
            self.counters[budget_key] = self.counters.get(budget_key, 0) + 1
            try:
                response = await self.http.request(
                    method,
                    url,
                    headers=headers,
                    json=json_body,
                    params=params,
                )
            except (httpx.TimeoutException, httpx.NetworkError) as exc:
                if attempt == 4:
                    raise PlacesRequestError(
                        0, "NETWORK_ERROR", type(exc).__name__
                    ) from exc
                await self._sleep(attempt)
                continue
            if response.is_success:
                await self._sleep()
                return response.json()
            try:
                payload = response.json().get("error") or {}
            except (ValueError, AttributeError):
                payload = {}
            status = str(payload.get("status") or f"HTTP_{response.status_code}")
            message = str(payload.get("message") or "Google Places request failed")
            last_error = PlacesRequestError(response.status_code, status, message)
            if (
                response.status_code not in {408, 429, 500, 502, 503, 504}
                or attempt == 4
            ):
                raise last_error
            await self._sleep(attempt)
        raise last_error or PlacesRequestError(0, "UNKNOWN", "Unknown request failure")

    async def resolve_center(
        self, state: str, district: str, taluka: str | None
    ) -> tuple[float, float, str] | None:
        target = taluka or district
        query = (
            f"{target}, {district}, {state}, India"
            if taluka
            else f"{district}, {state}, India"
        )
        payload = await self._request(
            "POST",
            PLACES_SEARCH_URL,
            field_mask=CENTER_FIELDS,
            operation="text",
            json_body={
                "textQuery": query,
                "pageSize": 5,
                "languageCode": "en",
                "regionCode": "IN",
            },
        )
        places = payload.get("places") or []
        best: tuple[int, dict[str, Any]] | None = None
        for place in places:
            lat, lng = place_coordinates(place)
            if lat is None or lng is None:
                continue
            found_state = canonical_state(
                component(place, "administrative_area_level_1")
            )
            found_district = canonical_district(
                component(place, "administrative_area_level_2")
            )
            score = 0
            if found_state == state:
                score += 8
            if found_district and normalize_dense(found_district) == normalize_dense(
                district
            ):
                score += 5
            if normalize_dense(target) in normalize_dense(display_name(place)):
                score += 3
            if normalize_dense(target) in normalize_dense(
                place.get("formattedAddress")
            ):
                score += 2
            if best is None or score > best[0]:
                best = (score, place)
        if not best:
            return None
        lat, lng = place_coordinates(best[1])
        if lat is None or lng is None:
            return None
        return lat, lng, "GOOGLE_PLACES_TALUKA" if taluka else "GOOGLE_PLACES_DISTRICT"

    async def search_rmc(
        self,
        query: str,
        center: tuple[float, float],
        radius_m: int,
    ) -> tuple[list[dict[str, Any]], int, bool]:
        page_token: str | None = None
        out: list[dict[str, Any]] = []
        pages = 0
        saturated = False
        while True:
            body: dict[str, Any] = {
                "textQuery": query,
                "pageSize": 20,
                "languageCode": "en",
                "regionCode": "IN",
                "locationBias": {
                    "circle": {
                        "center": {"latitude": center[0], "longitude": center[1]},
                        "radius": radius_m,
                    }
                },
            }
            if page_token:
                body["pageToken"] = page_token
            payload = await self._request(
                "POST",
                PLACES_SEARCH_URL,
                field_mask=SEARCH_FIELDS,
                operation="text",
                json_body=body,
            )
            pages += 1
            out.extend(payload.get("places") or [])
            page_token = payload.get("nextPageToken")
            if not page_token:
                break
            if pages >= self.config.max_pages:
                saturated = True
                break
            await asyncio.sleep(1.5)
        return out, pages, saturated

    async def details(self, place_id: str) -> dict[str, Any]:
        return await self._request(
            "GET",
            PLACES_DETAILS_URL.format(place_id=place_id),
            field_mask=DETAIL_FIELDS,
            operation="details",
            params={"languageCode": "en", "regionCode": "IN"},
        )


def load_manifest(path: Path, states: Iterable[str]) -> list[dict[str, str]]:
    wanted = set(states)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = [
            dict(row) for row in csv.DictReader(handle) if row.get("state") in wanted
        ]
    rows.sort(
        key=lambda row: (
            TARGET_STATES.index(row["state"]),
            row["district"],
            row["taluka"],
        )
    )
    seen: set[tuple[str, str, str]] = set()
    clean: list[dict[str, str]] = []
    for row in rows:
        key = tuple(
            normalize_dense(row[name]) for name in ("state", "district", "taluka")
        )
        if key in seen:
            continue
        seen.add(key)
        clean.append(row)
    return clean


def config_signature(config: Config, manifest: list[dict[str, str]]) -> str:
    payload = {
        "states": config.states,
        "terms": SEARCH_TERMS,
        "radius": config.search_radius_m,
        "grid_radius": config.grid_radius_m,
        "manifest": [
            (row["state"], row["district"], row["taluka"], row.get("source_code"))
            for row in manifest
        ],
    }
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()


class CheckpointStore:
    def __init__(self, config: Config, signature: str):
        self.config = config
        self.signature = signature
        self.local_path = config.local_output_dir / "checkpoint.json.gz"
        self.object_name = f"{config.output_prefix}/work-in-progress/checkpoint.json.gz"

    def _download(self) -> bool:
        if not self.config.output_bucket:
            return self.local_path.exists()
        try:
            from google.cloud import storage

            blob = (
                storage.Client()
                .bucket(self.config.output_bucket)
                .blob(self.object_name)
            )
            if not blob.exists():
                return False
            self.local_path.parent.mkdir(parents=True, exist_ok=True)
            blob.download_to_filename(str(self.local_path))
            return True
        except Exception as exc:
            print(
                json.dumps(
                    {"event": "checkpoint_download_failed", "error": type(exc).__name__}
                )
            )
            return False

    async def load(self) -> dict[str, Any] | None:
        if not self.config.resume:
            return None
        exists = await asyncio.to_thread(self._download)
        if not exists:
            return None
        try:
            with gzip.open(self.local_path, "rt", encoding="utf-8") as handle:
                data = json.load(handle)
            if data.get("signature") != self.signature:
                print(
                    json.dumps(
                        {"event": "checkpoint_ignored", "reason": "CONFIG_CHANGED"}
                    )
                )
                return None
            return data
        except Exception as exc:
            print(
                json.dumps(
                    {"event": "checkpoint_ignored", "reason": type(exc).__name__}
                )
            )
            return None

    def _upload(self) -> None:
        if not self.config.output_bucket:
            return
        from google.cloud import storage

        storage.Client().bucket(self.config.output_bucket).blob(
            self.object_name
        ).upload_from_filename(str(self.local_path), content_type="application/gzip")

    async def save(self, state: dict[str, Any]) -> None:
        self.local_path.parent.mkdir(parents=True, exist_ok=True)
        state["signature"] = self.signature
        state["checkpointed_at"] = datetime.now(timezone.utc).isoformat()
        temp_path = self.local_path.with_suffix(".tmp")
        with gzip.open(temp_path, "wt", encoding="utf-8") as handle:
            json.dump(state, handle, ensure_ascii=False, separators=(",", ":"))
        temp_path.replace(self.local_path)
        await asyncio.to_thread(self._upload)

    def clear(self) -> None:
        if self.local_path.exists():
            self.local_path.unlink()
        if self.config.output_bucket:
            try:
                from google.cloud import storage

                storage.Client().bucket(self.config.output_bucket).blob(
                    self.object_name
                ).delete()
            except Exception:
                pass


def initial_state() -> dict[str, Any]:
    return {
        "candidates": {},
        "details": {},
        "detail_errors": {},
        "coverage": [],
        "completed_units": [],
        "counters": {"text": 0, "details": 0},
        "budget_exhausted": None,
    }


def coverage_key(state: str, district: str, taluka: str) -> str:
    return f"{state}|{district}|{taluka}"


def merge_candidate(
    candidates: dict[str, dict[str, Any]],
    place: dict[str, Any],
    *,
    query: str,
    term: str,
    unit_key: str,
    state: str,
    district: str,
    taluka: str,
) -> str | None:
    place_id = str(place.get("id") or "").strip()
    if not place_id:
        return None
    target = candidates.setdefault(
        place_id,
        {
            "summary": place,
            "queries": [],
            "terms": [],
            "coverage_keys": [],
            "searched_states": [],
            "searched_districts": [],
            "searched_talukas": [],
        },
    )
    if len(json.dumps(place, ensure_ascii=False)) > len(
        json.dumps(target.get("summary") or {}, ensure_ascii=False)
    ):
        target["summary"] = place
    for key, value in (
        ("queries", query),
        ("terms", term),
        ("coverage_keys", unit_key),
        ("searched_states", state),
        ("searched_districts", district),
        ("searched_talukas", taluka),
    ):
        if value not in target[key]:
            target[key].append(value)
    return place_id


def grid_points(center: tuple[float, float]) -> list[tuple[float, float]]:
    lat, lng = center
    north_south = 0.18
    east_west = 0.18 / max(0.25, math.cos(math.radians(lat)))
    return [
        center,
        (lat + north_south, lng),
        (lat - north_south, lng),
        (lat, lng + east_west),
        (lat, lng - east_west),
    ]


async def find_center_with_fallback(
    client: PlacesClient,
    state: str,
    district: str,
    taluka: str,
    district_centers: dict[tuple[str, str], tuple[float, float, str] | None],
) -> tuple[float, float, str] | None:
    try:
        center = await client.resolve_center(state, district, taluka)
        if center:
            return center
    except (PlacesRequestError, BudgetExceeded):
        raise
    district_key = (state, district)
    if district_key not in district_centers:
        district_centers[district_key] = await client.resolve_center(
            state, district, None
        )
    fallback = district_centers[district_key]
    if fallback:
        return fallback[0], fallback[1], "DISTRICT_FALLBACK_LOCATION_REVIEW_REQUIRED"
    return None


async def collect_searches(
    config: Config,
    client: PlacesClient,
    manifest: list[dict[str, str]],
    state_data: dict[str, Any],
    checkpoint: CheckpointStore,
) -> None:
    candidates: dict[str, dict[str, Any]] = state_data["candidates"]
    completed = set(state_data["completed_units"])
    district_centers: dict[tuple[str, str], tuple[float, float, str] | None] = {}

    for index, row in enumerate(manifest, start=1):
        state, district, taluka = row["state"], row["district"], row["taluka"]
        unit_key = coverage_key(state, district, taluka)
        if unit_key in completed:
            continue
        coverage = {
            "state": state,
            "district": district,
            "taluka": taluka,
            "coverage_source": row.get("source") or "",
            "source_code": row.get("source_code") or "",
            "location_note": row.get("location_note") or "",
            "terms_executed": [],
            "raw_results": 0,
            "candidate_ids": [],
            "search_requests": 0,
            "center_latitude": None,
            "center_longitude": None,
            "center_source": "",
            "status": "IN_PROGRESS",
            "errors": [],
        }
        before_requests = client.counters["text"]
        try:
            center = await find_center_with_fallback(
                client, state, district, taluka, district_centers
            )
            if not center:
                coverage["status"] = "LOCATION_CENTER_FAILED"
                coverage["errors"].append(
                    "No reliable taluka or district center returned"
                )
            else:
                coverage["center_latitude"] = center[0]
                coverage["center_longitude"] = center[1]
                coverage["center_source"] = center[2]
                for term in SEARCH_TERMS:
                    query = f"{term} near {taluka}, {district}, {state}, India"
                    places, _pages, saturated = await client.search_rmc(
                        query, (center[0], center[1]), config.search_radius_m
                    )
                    coverage["terms_executed"].append(term)
                    coverage["raw_results"] += len(places)
                    for place in places:
                        place_id = merge_candidate(
                            candidates,
                            place,
                            query=query,
                            term=term,
                            unit_key=unit_key,
                            state=state,
                            district=district,
                            taluka=taluka,
                        )
                        if place_id and place_id not in coverage["candidate_ids"]:
                            coverage["candidate_ids"].append(place_id)
                    if saturated:
                        coverage["errors"].append(f"SATURATED_RESULTS:{term}")
                coverage["status"] = (
                    "LOCATION_REVIEW_REQUIRED"
                    if "LOCATION_REVIEW_REQUIRED" in coverage["center_source"]
                    else "COMPLETE"
                )
                if not coverage["candidate_ids"]:
                    coverage["status"] = "NO_RMC_RESULT_FOUND"
        except BudgetExceeded as exc:
            coverage["status"] = "INCOMPLETE_API_SAFETY_CAP"
            coverage["errors"].append(str(exc))
            state_data["budget_exhausted"] = str(exc)
        except PlacesRequestError as exc:
            coverage["status"] = "INCOMPLETE_API_ERROR"
            coverage["errors"].append(f"{exc.status}:{exc.safe_message}")
        coverage["search_requests"] = client.counters["text"] - before_requests
        state_data["coverage"].append(coverage)
        completed.add(unit_key)
        state_data["completed_units"] = sorted(completed)
        should_checkpoint = (
            index % config.checkpoint_every_units == 0
            or coverage["status"].startswith("INCOMPLETE")
            or index == len(manifest)
        )
        if should_checkpoint:
            await checkpoint.save(state_data)
        print(
            json.dumps(
                {
                    "event": "taluka_completed",
                    "progress": f"{index}/{len(manifest)}",
                    "state": state,
                    "district": district,
                    "taluka": taluka,
                    "status": coverage["status"],
                    "raw_count": coverage["raw_results"],
                    "unique_count": len(coverage["candidate_ids"]),
                    "text_requests": client.counters["text"],
                },
                ensure_ascii=False,
            )
        )
        if state_data.get("budget_exhausted"):
            return

    await collect_adaptive_district_grids(
        config, client, manifest, state_data, checkpoint, district_centers
    )


async def collect_adaptive_district_grids(
    config: Config,
    client: PlacesClient,
    manifest: list[dict[str, str]],
    state_data: dict[str, Any],
    checkpoint: CheckpointStore,
    district_centers: dict[tuple[str, str], tuple[float, float, str] | None],
) -> None:
    candidates = state_data["candidates"]
    completed = set(state_data["completed_units"])
    districts = sorted(
        {(row["state"], row["district"]) for row in manifest},
        key=lambda pair: (TARGET_STATES.index(pair[0]), pair[1]),
    )
    for grid_index, (state, district) in enumerate(districts, start=1):
        if (state, district) not in HIGH_DENSITY_DISTRICTS:
            continue
        taluka = "[DISTRICT GRID]"
        unit_key = coverage_key(state, district, taluka)
        if unit_key in completed:
            continue
        coverage = {
            "state": state,
            "district": district,
            "taluka": taluka,
            "coverage_source": "Adaptive high-density grid",
            "source_code": "",
            "location_note": "DISTRICT_GRID",
            "terms_executed": [],
            "raw_results": 0,
            "candidate_ids": [],
            "search_requests": 0,
            "center_latitude": None,
            "center_longitude": None,
            "center_source": "GOOGLE_PLACES_DISTRICT_GRID",
            "status": "IN_PROGRESS",
            "errors": [],
        }
        before = client.counters["text"]
        try:
            center = district_centers.get((state, district))
            if center is None:
                center = await client.resolve_center(state, district, None)
                district_centers[(state, district)] = center
            if not center:
                coverage["status"] = "LOCATION_CENTER_FAILED"
            else:
                coverage["center_latitude"], coverage["center_longitude"] = center[:2]
                for point_index, point in enumerate(
                    grid_points((center[0], center[1]))
                ):
                    for term in SEARCH_TERMS:
                        query = f"{term} in {district}, {state}, India"
                        places, _pages, saturated = await client.search_rmc(
                            query, point, config.grid_radius_m
                        )
                        executed = f"{term} [grid {point_index + 1}]"
                        coverage["terms_executed"].append(executed)
                        coverage["raw_results"] += len(places)
                        for place in places:
                            place_id = merge_candidate(
                                candidates,
                                place,
                                query=f"{query} [grid {point_index + 1}]",
                                term=term,
                                unit_key=unit_key,
                                state=state,
                                district=district,
                                taluka=taluka,
                            )
                            if place_id and place_id not in coverage["candidate_ids"]:
                                coverage["candidate_ids"].append(place_id)
                        if saturated:
                            coverage["errors"].append(f"SATURATED_RESULTS:{executed}")
                coverage["status"] = (
                    "COMPLETE" if coverage["candidate_ids"] else "NO_RMC_RESULT_FOUND"
                )
        except BudgetExceeded as exc:
            coverage["status"] = "INCOMPLETE_API_SAFETY_CAP"
            coverage["errors"].append(str(exc))
            state_data["budget_exhausted"] = str(exc)
        except PlacesRequestError as exc:
            coverage["status"] = "INCOMPLETE_API_ERROR"
            coverage["errors"].append(f"{exc.status}:{exc.safe_message}")
        coverage["search_requests"] = client.counters["text"] - before
        state_data["coverage"].append(coverage)
        completed.add(unit_key)
        state_data["completed_units"] = sorted(completed)
        if grid_index % config.checkpoint_every_units == 0 or coverage[
            "status"
        ].startswith("INCOMPLETE"):
            await checkpoint.save(state_data)
        print(
            json.dumps(
                {
                    "event": "district_grid_completed",
                    "state": state,
                    "district": district,
                    "status": coverage["status"],
                    "raw_count": coverage["raw_results"],
                    "unique_count": len(coverage["candidate_ids"]),
                    "text_requests": client.counters["text"],
                }
            )
        )
        if state_data.get("budget_exhausted"):
            return
    await checkpoint.save(state_data)


async def collect_details(
    client: PlacesClient,
    state_data: dict[str, Any],
    checkpoint: CheckpointStore,
) -> None:
    place_ids = sorted(state_data["candidates"])
    for index, place_id in enumerate(place_ids, start=1):
        if place_id in state_data["details"] or place_id in state_data["detail_errors"]:
            continue
        try:
            state_data["details"][place_id] = await client.details(place_id)
        except BudgetExceeded as exc:
            state_data["budget_exhausted"] = str(exc)
            break
        except PlacesRequestError as exc:
            state_data["detail_errors"][place_id] = {
                "status": exc.status,
                "message": exc.safe_message,
            }
        if index % 100 == 0:
            await checkpoint.save(state_data)
            print(
                json.dumps(
                    {
                        "event": "details_progress",
                        "processed": index,
                        "total": len(place_ids),
                        "details_requests": client.counters["details"],
                    }
                )
            )
    await checkpoint.save(state_data)


def classify_relevance(place: dict[str, Any]) -> tuple[str, str]:
    primary_display = place.get("primaryTypeDisplayName") or {}
    text = " ".join(
        str(value or "")
        for value in (
            display_name(place),
            place.get("formattedAddress"),
            place.get("primaryType"),
            primary_display.get("text"),
            " ".join(place.get("types") or []),
        )
    )
    strong = bool(STRONG_RMC_RE.search(text))
    false_positive = bool(FALSE_POSITIVE_RE.search(text))
    concrete = bool(CONCRETE_CUE_RE.search(text))
    if false_positive and not strong:
        return "EXCLUDE", "OBVIOUS_FALSE_POSITIVE"
    if strong:
        if re.search(r"supplier|dealer", text, re.IGNORECASE):
            return "RMC Supplier", "STRONG"
        return "RMC Plant", "STRONG"
    if concrete:
        return "Concrete Supplier", "REVIEW"
    return "EXCLUDE", "NO_RMC_RELEVANCE_SIGNAL"


def known_districts(manifest: list[dict[str, str]]) -> dict[str, set[str]]:
    result: dict[str, set[str]] = defaultdict(set)
    for row in manifest:
        result[row["state"]].add(normalize_dense(row["district"]))
    return result


def known_talukas(manifest: list[dict[str, str]]) -> dict[tuple[str, str], set[str]]:
    result: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in manifest:
        result[(row["state"], normalize_dense(row["district"]))].add(
            normalize_dense(row["taluka"])
        )
    return result


def build_record(
    place_id: str,
    candidate: dict[str, Any],
    detail: dict[str, Any] | None,
    detail_error: dict[str, str] | None,
    manifest_districts: dict[str, set[str]],
    manifest_talukas: dict[tuple[str, str], set[str]],
) -> tuple[dict[str, Any] | None, str]:
    place = detail or candidate.get("summary") or {}
    name = display_name(place)
    lat, lng = place_coordinates(place)
    if not place_id or not name or lat is None or lng is None:
        return None, "INCOMPLETE_PLACE_CORE_FIELDS"
    record_type, relevance = classify_relevance(place)
    if record_type == "EXCLUDE":
        return None, relevance

    state = canonical_state(component(place, "administrative_area_level_1"))
    searched_states = candidate.get("searched_states") or []
    if state is None and len(set(searched_states)) == 1:
        state = searched_states[0]
        location_issue = True
    else:
        location_issue = False
    if state not in TARGET_STATES:
        return None, "OUTSIDE_TARGET_STATE"
    if searched_states and state not in searched_states:
        return None, "OUTSIDE_SEARCHED_STATE"

    district = canonical_district(component(place, "administrative_area_level_2"))
    taluka = canonical_taluka(component(place, "administrative_area_level_3"))
    city = component(place, "locality", "postal_town", "sublocality_level_1")
    pin_code = component(place, "postal_code")
    if not district or normalize_dense(district) not in manifest_districts[state]:
        location_issue = True
    if not taluka:
        location_issue = True
    elif district:
        district_talukas = manifest_talukas.get(
            (state, normalize_dense(district)), set()
        )
        if district_talukas and normalize_dense(taluka) not in district_talukas:
            location_issue = True

    phone = str(place.get("nationalPhoneNumber") or "").strip()
    international_phone = str(place.get("internationalPhoneNumber") or "").strip()
    status = str(place.get("businessStatus") or "").strip()
    closed = status in {"CLOSED_PERMANENTLY", "CLOSED_TEMPORARILY"}
    review_required = (
        relevance == "REVIEW" or location_issue or detail_error is not None or closed
    )
    if review_required:
        quality = "C"
        verification = (
            "LOCATION_REVIEW_REQUIRED" if location_issue else "REVIEW_REQUIRED"
        )
    elif phone:
        quality = "A"
        verification = "GOOGLE_PLACES_VERIFIED"
    else:
        quality = "B"
        verification = "GOOGLE_PLACES_VERIFIED"

    maps_url = str(place.get("googleMapsUri") or "").strip()
    if not maps_url.startswith("https://"):
        maps_url = f"https://www.google.com/maps/place/?q=place_id:{place_id}"
    primary_type = str(place.get("primaryType") or "").strip()
    queries = candidate.get("queries") or []
    terms = candidate.get("terms") or []
    address_components = place.get("addressComponents") or []
    return (
        {
            "Plant / Supplier Name": name,
            "Record Type": record_type,
            "Public Business Phone": phone or None,
            "International Phone": international_phone or None,
            "Full Address": str(place.get("formattedAddress") or "").strip() or None,
            "State": state,
            "District": district,
            "Taluka / Tehsil": taluka,
            "City / Locality": city,
            "PIN Code": pin_code,
            "Latitude": lat,
            "Longitude": lng,
            "Google Place ID": place_id,
            "Google Maps URL": maps_url,
            "Website": str(place.get("websiteUri") or "").strip() or None,
            "Rating": place.get("rating"),
            "Review Count": place.get("userRatingCount"),
            "Business Status": status or None,
            "Google Primary Type": primary_type or None,
            "Search Term": terms[0] if terms else None,
            "Data Quality": quality,
            "Verification Status": verification,
            "Last Verified Date": date.today(),
            "Phone Status": "PHONE_PUBLIC" if phone else "PHONE_NOT_PUBLIC",
            "Discovery Queries": " | ".join(queries),
            "Address Components JSON": json.dumps(
                address_components, ensure_ascii=False, separators=(",", ":")
            ),
            "Searched Districts": " | ".join(candidate.get("searched_districts") or []),
            "Searched Talukas": " | ".join(candidate.get("searched_talukas") or []),
            "_coverage_keys": list(candidate.get("coverage_keys") or []),
            "_alias_place_ids": [],
        },
        "RETAINED",
    )


def record_richness(record: dict[str, Any]) -> int:
    return sum(
        1
        for key, value in record.items()
        if not key.startswith("_") and value not in (None, "")
    )


def merge_records(canonical: dict[str, Any], duplicate: dict[str, Any]) -> None:
    if record_richness(duplicate) > record_richness(canonical):
        preserve = {
            "_coverage_keys": canonical.get("_coverage_keys", []),
            "_alias_place_ids": canonical.get("_alias_place_ids", []),
        }
        original_id = canonical["Google Place ID"]
        canonical.clear()
        canonical.update(duplicate)
        canonical["_alias_place_ids"] = preserve["_alias_place_ids"] + [original_id]
        canonical["_coverage_keys"] = preserve["_coverage_keys"] + canonical.get(
            "_coverage_keys", []
        )
    else:
        canonical["_alias_place_ids"].append(duplicate["Google Place ID"])
        canonical["_coverage_keys"].extend(duplicate.get("_coverage_keys", []))
    canonical["_alias_place_ids"] = sorted(set(canonical["_alias_place_ids"]))
    canonical["_coverage_keys"] = sorted(set(canonical["_coverage_keys"]))
    for key in ("Discovery Queries", "Searched Districts", "Searched Talukas"):
        values = []
        for value in (canonical.get(key), duplicate.get(key)):
            for item in str(value or "").split(" | "):
                if item and item not in values:
                    values.append(item)
        canonical[key] = " | ".join(values)


def should_merge_records(left: dict[str, Any], right: dict[str, Any]) -> bool:
    if left["Google Place ID"] == right["Google Place ID"]:
        return True
    left_coord = (float(left["Latitude"]), float(left["Longitude"]))
    right_coord = (float(right["Latitude"]), float(right["Longitude"]))
    distance = haversine_m(left_coord, right_coord)
    name_similarity = SequenceMatcher(
        None,
        normalize_text(left["Plant / Supplier Name"]),
        normalize_text(right["Plant / Supplier Name"]),
    ).ratio()
    address_similarity = SequenceMatcher(
        None,
        normalize_text(left.get("Full Address")),
        normalize_text(right.get("Full Address")),
    ).ratio()
    same_phone = bool(
        normalize_phone(left.get("Public Business Phone"))
        and normalize_phone(left.get("Public Business Phone"))
        == normalize_phone(right.get("Public Business Phone"))
    )
    if distance <= 60 and name_similarity >= 0.78:
        return True
    if name_similarity >= 0.93 and address_similarity >= 0.88:
        return True
    # A shared corporate phone must not collapse physically separate plants.
    if same_phone and (distance <= 1_000 or address_similarity >= 0.90):
        return True
    return False


def deduplicate_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    canonical: list[dict[str, Any]] = []
    for record in sorted(
        records,
        key=lambda row: (
            TARGET_STATES.index(row["State"]),
            row.get("District") or "",
            row.get("Taluka / Tehsil") or "",
            row["Plant / Supplier Name"],
        ),
    ):
        match = next(
            (row for row in canonical if should_merge_records(row, record)), None
        )
        if match:
            merge_records(match, record)
        else:
            canonical.append(record)
    return canonical


def produce_records(
    manifest: list[dict[str, str]], state_data: dict[str, Any]
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    districts = known_districts(manifest)
    talukas = known_talukas(manifest)
    records: list[dict[str, Any]] = []
    exclusions: dict[str, str] = {}
    for place_id, candidate in state_data["candidates"].items():
        detail = state_data["details"].get(place_id)
        detail_error = state_data["detail_errors"].get(place_id)
        record, reason = build_record(
            place_id, candidate, detail, detail_error, districts, talukas
        )
        if record:
            records.append(record)
        else:
            exclusions[place_id] = reason
    return deduplicate_records(records), exclusions


def finalize_coverage(
    coverage: list[dict[str, Any]],
    records: list[dict[str, Any]],
    budget_exhausted: str | None,
) -> list[dict[str, Any]]:
    retained_ids: set[str] = set()
    for record in records:
        retained_ids.add(record["Google Place ID"])
        retained_ids.update(record.get("_alias_place_ids") or [])
    out = []
    for item in coverage:
        ids = set(item.get("candidate_ids") or [])
        retained = len(ids & retained_ids)
        status = item.get("status") or "UNKNOWN"
        if budget_exhausted and status in {"IN_PROGRESS", "COMPLETE"}:
            status = "INCOMPLETE_API_SAFETY_CAP"
        out.append(
            {
                "State": item.get("state"),
                "District": item.get("district"),
                "Taluka": item.get("taluka"),
                "Coverage Source": item.get("coverage_source"),
                "Coverage Source URL": COVERAGE_SOURCE_URLS.get(
                    item.get("coverage_source") or "", ""
                ),
                "Source Code": item.get("source_code"),
                "Location Note": item.get("location_note"),
                "Search Terms Executed": " | ".join(item.get("terms_executed") or []),
                "Raw Results": item.get("raw_results", 0),
                "Unique Results": len(ids),
                "Retained": retained,
                "Excluded": max(0, len(ids) - retained),
                "Search Requests": item.get("search_requests", 0),
                "Center Latitude": item.get("center_latitude"),
                "Center Longitude": item.get("center_longitude"),
                "Center Source": item.get("center_source"),
                "Search Completion Status": status,
                "Errors": " | ".join(item.get("errors") or []),
            }
        )
    return out


def clean_record(record: dict[str, Any], serial: int) -> dict[str, Any]:
    return {
        "Sr No": serial,
        **{key: record.get(key) for key in MASTER_COLUMNS if key != "Sr No"},
    }


def summarize_by(
    records: list[dict[str, Any]], group_fields: tuple[str, ...]
) -> list[dict[str, Any]]:
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        groups[
            tuple(record.get(field) or "UNASSIGNED" for field in group_fields)
        ].append(record)
    out: list[dict[str, Any]] = []
    for key in sorted(groups):
        rows = groups[key]
        item = {field: value for field, value in zip(group_fields, key)}
        item.update(
            {
                "Total Records": len(rows),
                "With Phone": sum(
                    bool(row.get("Public Business Phone")) for row in rows
                ),
                "Missing Phone": sum(
                    not bool(row.get("Public Business Phone")) for row in rows
                ),
                "Review Required": sum(row.get("Data Quality") == "C" for row in rows),
                "Quality A": sum(row.get("Data Quality") == "A" for row in rows),
                "Quality B": sum(row.get("Data Quality") == "B" for row in rows),
                "Quality C": sum(row.get("Data Quality") == "C" for row in rows),
            }
        )
        out.append(item)
    return out


HEADER_FILL = PatternFill("solid", fgColor="08765F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTLE_BORDER = Border(bottom=Side(style="thin", color="D6E1DD"))
QUALITY_FILLS = {
    "A": PatternFill("solid", fgColor="D9EAD3"),
    "B": PatternFill("solid", fgColor="FFF2CC"),
    "C": PatternFill("solid", fgColor="FCE5CD"),
}


def add_table_sheet(
    workbook: Workbook,
    name: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    table_name: str,
) -> Any:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    header = sheet[1]
    for cell in header:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    sheet.row_dimensions[1].height = 34
    if rows:
        reference = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    return sheet


def style_master_sheet(sheet: Any, columns: list[str], row_count: int) -> None:
    index = {column: i + 1 for i, column in enumerate(columns)}
    date_col = index.get("Last Verified Date")
    lat_col = index.get("Latitude")
    lng_col = index.get("Longitude")
    rating_col = index.get("Rating")
    review_col = index.get("Review Count")
    quality_col = index.get("Data Quality")
    maps_col = index.get("Google Maps URL")
    website_col = index.get("Website")
    for row in range(2, row_count + 2):
        if date_col:
            sheet.cell(row, date_col).number_format = "yyyy-mm-dd"
        for col in (lat_col, lng_col):
            if col:
                sheet.cell(row, col).number_format = "0.000000"
        if rating_col:
            sheet.cell(row, rating_col).number_format = "0.0"
        if review_col:
            sheet.cell(row, review_col).number_format = "#,##0"
        if quality_col:
            cell = sheet.cell(row, quality_col)
            if cell.value in QUALITY_FILLS:
                cell.fill = QUALITY_FILLS[cell.value]
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
        for col in (maps_col, website_col):
            if col:
                cell = sheet.cell(row, col)
                if isinstance(cell.value, str) and cell.value.startswith("https://"):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
    preferred = {
        "Sr No": 9,
        "Plant / Supplier Name": 34,
        "Record Type": 20,
        "Public Business Phone": 20,
        "International Phone": 21,
        "Full Address": 52,
        "State": 16,
        "District": 26,
        "Taluka / Tehsil": 25,
        "City / Locality": 25,
        "PIN Code": 12,
        "Latitude": 14,
        "Longitude": 14,
        "Google Place ID": 34,
        "Google Maps URL": 44,
        "Website": 36,
        "Rating": 10,
        "Review Count": 13,
        "Business Status": 20,
        "Google Primary Type": 24,
        "Search Term": 28,
        "Data Quality": 12,
        "Verification Status": 27,
        "Last Verified Date": 18,
        "Phone Status": 20,
        "Discovery Queries": 60,
        "Address Components JSON": 60,
        "Searched Districts": 34,
        "Searched Talukas": 34,
    }
    for column, width in preferred.items():
        if column in index:
            sheet.column_dimensions[get_column_letter(index[column])].width = width
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(columns))}{max(1, row_count + 1)}"
    )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.border = SUBTLE_BORDER


def create_workbook(
    state: str,
    records: list[dict[str, Any]],
    coverage: list[dict[str, Any]],
    output_path: Path,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    clean = [
        clean_record(record, index) for index, record in enumerate(records, start=1)
    ]
    master = add_table_sheet(workbook, "MASTER", MASTER_COLUMNS, clean, "MasterRecords")
    style_master_sheet(master, MASTER_COLUMNS, len(clean))

    district_rows = summarize_by(records, ("District",))
    district_columns = [
        "District",
        "Total Records",
        "With Phone",
        "Missing Phone",
        "Review Required",
        "Quality A",
        "Quality B",
        "Quality C",
    ]
    district = add_table_sheet(
        workbook, "BY_DISTRICT", district_columns, district_rows, "DistrictSummary"
    )

    taluka_rows = summarize_by(records, ("District", "Taluka / Tehsil"))
    taluka_columns = [
        "District",
        "Taluka / Tehsil",
        "Total Records",
        "With Phone",
        "Missing Phone",
        "Review Required",
        "Quality A",
        "Quality B",
        "Quality C",
    ]
    taluka = add_table_sheet(
        workbook, "BY_TALUKA", taluka_columns, taluka_rows, "TalukaSummary"
    )

    missing_rows = [row for row in clean if row["Phone Status"] == "PHONE_NOT_PUBLIC"]
    missing = add_table_sheet(
        workbook, "MISSING_PHONE", MASTER_COLUMNS, missing_rows, "MissingPhoneRecords"
    )
    style_master_sheet(missing, MASTER_COLUMNS, len(missing_rows))

    review_rows = [row for row in clean if row["Data Quality"] == "C"]
    review = add_table_sheet(
        workbook, "REVIEW_REQUIRED", MASTER_COLUMNS, review_rows, "ReviewRecords"
    )
    style_master_sheet(review, MASTER_COLUMNS, len(review_rows))

    state_coverage = [row for row in coverage if row["State"] == state]
    coverage_sheet = add_table_sheet(
        workbook,
        "SEARCH_COVERAGE",
        COVERAGE_COLUMNS,
        state_coverage,
        "SearchCoverage",
    )
    for summary_sheet in (district, taluka, coverage_sheet):
        summary_sheet.freeze_panes = "A2"
        for column_cells in summary_sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            summary_sheet.column_dimensions[letter].width = min(
                55, max(12, max_length + 2)
            )
        for row in summary_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                cell.border = SUBTLE_BORDER

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    # A second parser pass catches corrupt XLSX packages before upload.
    verified = load_workbook(output_path, read_only=True, data_only=False)
    expected = {
        "MASTER",
        "BY_DISTRICT",
        "BY_TALUKA",
        "MISSING_PHONE",
        "REVIEW_REQUIRED",
        "SEARCH_COVERAGE",
    }
    if set(verified.sheetnames) != expected:
        raise RuntimeError(f"Workbook verification failed for {state}")
    verified.close()


def state_filename(state: str) -> str:
    return f"{state.upper()}_RMC_PLACES_MASTER.xlsx"


def build_report(
    config: Config,
    manifest: list[dict[str, str]],
    records: list[dict[str, Any]],
    coverage: list[dict[str, Any]],
    state_data: dict[str, Any],
) -> dict[str, Any]:
    states: list[dict[str, Any]] = []
    for state in config.states:
        state_manifest = [row for row in manifest if row["state"] == state]
        state_records = [row for row in records if row["State"] == state]
        state_coverage = [
            row
            for row in coverage
            if row["State"] == state and row["Taluka"] != "[DISTRICT GRID]"
        ]
        states.append(
            {
                "state": state,
                "districts_planned": len({row["district"] for row in state_manifest}),
                "districts_searched": len(
                    {
                        row["District"]
                        for row in state_coverage
                        if row["Search Terms Executed"]
                    }
                ),
                "talukas_planned": len(state_manifest),
                "talukas_searched": sum(
                    bool(row["Search Terms Executed"]) for row in state_coverage
                ),
                "raw_places_results": sum(
                    int(row["Raw Results"] or 0) for row in state_coverage
                ),
                "unique_rmc_records": len(state_records),
                "with_phone": sum(
                    bool(row.get("Public Business Phone")) for row in state_records
                ),
                "missing_phone": sum(
                    not bool(row.get("Public Business Phone")) for row in state_records
                ),
                "review_required": sum(
                    row.get("Data Quality") == "C" for row in state_records
                ),
                "no_result_talukas": [
                    f"{row['District']} / {row['Taluka']}"
                    for row in state_coverage
                    if row["Search Completion Status"] == "NO_RMC_RESULT_FOUND"
                ],
                "weak_or_incomplete_coverage": [
                    f"{row['District']} / {row['Taluka']}: {row['Search Completion Status']}"
                    for row in state_coverage
                    if row["Search Completion Status"]
                    not in {"COMPLETE", "NO_RMC_RESULT_FOUND"}
                ],
            }
        )
    return {
        "run_id": datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "complete": (
            not state_data.get("budget_exhausted")
            and not state_data["detail_errors"]
            and len(state_data["details"]) == len(state_data["candidates"])
            and all(
                row["Search Completion Status"] in {"COMPLETE", "NO_RMC_RESULT_FOUND"}
                for row in coverage
            )
        ),
        "api_safety_cap_or_quota_issue": state_data.get("budget_exhausted"),
        "text_search_requests": state_data["counters"].get("text", 0),
        "place_details_requests": state_data["counters"].get("details", 0),
        "candidate_place_ids": len(state_data["candidates"]),
        "place_details_retrieved": len(state_data["details"]),
        "place_details_failed": len(state_data["detail_errors"]),
        "states": states,
        "coverage_sources": sorted({row.get("source") or "" for row in manifest}),
        "security": {
            "secret_value_logged": False,
            "production_database_used": False,
            "production_service_modified": False,
        },
    }


def write_json_gzip(path: Path, payload: Any) -> None:
    with gzip.open(path, "wt", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))


def upload_outputs(config: Config, files: list[Path], run_id: str) -> None:
    if not config.output_bucket:
        return
    from google.cloud import storage

    bucket = storage.Client().bucket(config.output_bucket)
    for path in files:
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if path.suffix == ".xlsx"
            else "application/gzip" if path.suffix == ".gz" else "application/json"
        )
        for object_name in (
            f"{config.output_prefix}/runs/{run_id}/{path.name}",
            f"{config.output_prefix}/latest/{path.name}",
        ):
            bucket.blob(object_name).upload_from_filename(
                str(path), content_type=content_type
            )


def validate_outputs(
    config: Config,
    records: list[dict[str, Any]],
    coverage: list[dict[str, Any]],
    files: list[Path],
) -> None:
    place_ids = [record["Google Place ID"] for record in records]
    if len(place_ids) != len(set(place_ids)):
        raise RuntimeError("Duplicate canonical Place IDs remain")
    if any(record["State"] not in config.states for record in records):
        raise RuntimeError("A record belongs to an unrequested state")
    if any("bengaluru" in normalize_text(record.get("State")) for record in records):
        raise RuntimeError("Bengaluru was incorrectly classified as a state")
    if any(
        not str(record.get("Google Maps URL") or "").startswith("https://")
        for record in records
    ):
        raise RuntimeError("A record has an invalid Maps URL")
    if any(not path.exists() or path.stat().st_size == 0 for path in files):
        raise RuntimeError("An output file is missing or empty")
    covered_states = {row["State"] for row in coverage}
    if covered_states != set(config.states):
        raise RuntimeError("Search coverage is missing a requested state")


async def run() -> int:
    config = Config.from_env()
    config.validate()
    root = Path(__file__).resolve().parent
    manifest = load_manifest(root / "data" / "coverage.csv", config.states)
    signature = config_signature(config, manifest)
    config.local_output_dir.mkdir(parents=True, exist_ok=True)
    checkpoint = CheckpointStore(config, signature)
    state_data = await checkpoint.load() or initial_state()
    state_data.setdefault("counters", {"text": 0, "details": 0})
    print(
        json.dumps(
            {
                "event": "collector_started",
                "states": config.states,
                "districts": len({(r["state"], r["district"]) for r in manifest}),
                "talukas": len(manifest),
                "max_text_requests": config.max_text_requests,
                "max_detail_requests": config.max_detail_requests,
                "resumed": bool(state_data["completed_units"]),
            }
        )
    )
    client = PlacesClient(config, state_data["counters"])
    try:
        await collect_searches(config, client, manifest, state_data, checkpoint)
        await collect_details(client, state_data, checkpoint)
    finally:
        await client.close()

    records, exclusions = produce_records(manifest, state_data)
    coverage = finalize_coverage(
        state_data["coverage"], records, state_data.get("budget_exhausted")
    )
    output_files: list[Path] = []
    for state in config.states:
        path = config.local_output_dir / state_filename(state)
        create_workbook(
            state,
            [record for record in records if record["State"] == state],
            coverage,
            path,
        )
        output_files.append(path)

    report = build_report(config, manifest, records, coverage, state_data)
    run_id = report["run_id"]
    report_path = config.local_output_dir / "coverage-report.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    raw_path = config.local_output_dir / "raw-candidate-audit.json.gz"
    write_json_gzip(
        raw_path,
        {
            "candidates": state_data["candidates"],
            "details": state_data["details"],
            "detail_errors": state_data["detail_errors"],
            "exclusions": exclusions,
        },
    )
    output_files.extend((report_path, raw_path))
    validate_outputs(config, records, coverage, output_files)
    await asyncio.to_thread(upload_outputs, config, output_files, run_id)
    if report["complete"]:
        await asyncio.to_thread(checkpoint.clear)
    print(
        json.dumps(
            {
                "event": "collector_finished",
                "run_id": run_id,
                "complete": report["complete"],
                "records": len(records),
                "with_phone": sum(
                    bool(record.get("Public Business Phone")) for record in records
                ),
                "review_required": sum(
                    record.get("Data Quality") == "C" for record in records
                ),
                "output_files": [path.name for path in output_files],
            }
        )
    )
    return 0 if report["complete"] else 2


def main() -> None:
    try:
        raise SystemExit(asyncio.run(run()))
    except KeyboardInterrupt:
        raise SystemExit(130) from None
    except Exception as exc:
        # Never stringify arbitrary request objects: they could contain headers.
        print(
            json.dumps(
                {
                    "event": "collector_failed",
                    "error_type": type(exc).__name__,
                    "error": str(exc)[:300],
                }
            ),
            file=sys.stderr,
        )
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
