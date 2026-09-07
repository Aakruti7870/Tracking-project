from pathlib import Path

from openpyxl import load_workbook

import collector

ROOT = Path(__file__).resolve().parents[1]


def place(
    name="Alpha RMC Plant", place_id="pid-1", lat=19.0, lng=73.0, phone="9876543210"
):
    return {
        "id": place_id,
        "displayName": {"text": name},
        "formattedAddress": "Panvel, Raigad, Maharashtra 410206, India",
        "addressComponents": [
            {"longText": "Maharashtra", "types": ["administrative_area_level_1"]},
            {"longText": "Raigad", "types": ["administrative_area_level_2"]},
            {"longText": "Panvel", "types": ["administrative_area_level_3"]},
            {"longText": "Panvel", "types": ["locality"]},
            {"longText": "410206", "types": ["postal_code"]},
        ],
        "location": {"latitude": lat, "longitude": lng},
        "nationalPhoneNumber": phone,
        "internationalPhoneNumber": "+91 98765 43210" if phone else None,
        "googleMapsUri": f"https://maps.google.com/?cid={place_id}",
        "websiteUri": "https://example.com",
        "rating": 4.2,
        "userRatingCount": 18,
        "businessStatus": "OPERATIONAL",
        "primaryType": "concrete_contractor",
        "types": ["concrete_contractor"],
    }


def candidate():
    return {
        "summary": {},
        "queries": ["RMC Plant near Panvel"],
        "terms": ["RMC Plant"],
        "coverage_keys": ["Maharashtra|Raigad|Panvel"],
        "searched_states": ["Maharashtra"],
        "searched_districts": ["Raigad"],
        "searched_talukas": ["Panvel"],
    }


def test_manifest_has_four_separate_states_and_bengaluru_in_karnataka():
    rows = collector.load_manifest(
        ROOT / "data" / "coverage.csv", collector.TARGET_STATES
    )
    counts = {
        state: sum(row["state"] == state for row in rows)
        for state in collector.TARGET_STATES
    }
    assert counts == {"Maharashtra": 359, "Goa": 12, "Karnataka": 239, "Gujarat": 270}
    assert all(
        row["state"] == "Karnataka"
        for row in rows
        if "bengaluru" in row["district"].lower()
    )
    assert {row["district"] for row in rows if row["state"] == "Karnataka"} >= {
        "Bengaluru Urban",
        "Bengaluru Rural",
        "Bengaluru South",
    }


def test_false_positive_is_excluded_without_rmc_signal():
    record = place(name="Shree Cement and Hardware Shop")
    assert collector.classify_relevance(record) == ("EXCLUDE", "OBVIOUS_FALSE_POSITIVE")


def test_strong_rmc_signal_survives_generic_primary_type():
    record = place(name="ABC Ready Mix Concrete Plant")
    assert collector.classify_relevance(record)[0] == "RMC Plant"


def test_missing_phone_is_blank_and_marked_phone_not_public():
    manifest = [{"state": "Maharashtra", "district": "Raigad", "taluka": "Panvel"}]
    record, reason = collector.build_record(
        "pid-1",
        candidate(),
        place(phone=""),
        None,
        collector.known_districts(manifest),
        collector.known_talukas(manifest),
    )
    assert reason == "RETAINED"
    assert record["Public Business Phone"] is None
    assert record["Phone Status"] == "PHONE_NOT_PUBLIC"
    assert record["Data Quality"] == "B"


def test_shared_phone_does_not_merge_separate_physical_plants():
    manifest = [{"state": "Maharashtra", "district": "Raigad", "taluka": "Panvel"}]
    districts = collector.known_districts(manifest)
    talukas = collector.known_talukas(manifest)
    first, _ = collector.build_record(
        "pid-1", candidate(), place(), None, districts, talukas
    )
    second_place = place(place_id="pid-2", lat=19.5, lng=73.5)
    second_place["formattedAddress"] = "Different physical plant, Maharashtra, India"
    second, _ = collector.build_record(
        "pid-2", candidate(), second_place, None, districts, talukas
    )
    assert len(collector.deduplicate_records([first, second])) == 2


def test_near_duplicate_name_and_location_merges():
    manifest = [{"state": "Maharashtra", "district": "Raigad", "taluka": "Panvel"}]
    districts = collector.known_districts(manifest)
    talukas = collector.known_talukas(manifest)
    first, _ = collector.build_record(
        "pid-1", candidate(), place(), None, districts, talukas
    )
    duplicate = place(
        name="Alpha RMC Plants", place_id="pid-2", lat=19.0001, lng=73.0001
    )
    second, _ = collector.build_record(
        "pid-2", candidate(), duplicate, None, districts, talukas
    )
    result = collector.deduplicate_records([first, second])
    assert len(result) == 1
    assert len(result[0]["_alias_place_ids"]) == 1


def test_workbook_has_required_sheets_and_columns(tmp_path):
    manifest = [{"state": "Maharashtra", "district": "Raigad", "taluka": "Panvel"}]
    record, _ = collector.build_record(
        "pid-1",
        candidate(),
        place(),
        None,
        collector.known_districts(manifest),
        collector.known_talukas(manifest),
    )
    coverage = [
        {
            "State": "Maharashtra",
            "District": "Raigad",
            "Taluka": "Panvel",
            "Coverage Source": "test",
            "Source Code": "1",
            "Search Terms Executed": "RMC Plant",
            "Raw Results": 1,
            "Unique Results": 1,
            "Retained": 1,
            "Excluded": 0,
            "Search Requests": 1,
            "Center Latitude": 19.0,
            "Center Longitude": 73.0,
            "Center Source": "test",
            "Search Completion Status": "COMPLETE",
            "Errors": "",
        }
    ]
    output = tmp_path / "test.xlsx"
    collector.create_workbook("Maharashtra", [record], coverage, output)
    workbook = load_workbook(output, read_only=False)
    assert workbook.sheetnames == [
        "MASTER",
        "BY_DISTRICT",
        "BY_TALUKA",
        "MISSING_PHONE",
        "REVIEW_REQUIRED",
        "SEARCH_COVERAGE",
    ]
    assert [cell.value for cell in workbook["MASTER"][1]] == collector.MASTER_COLUMNS
    assert workbook["MASTER"]["O2"].hyperlink is not None
